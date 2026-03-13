import pandas as pd
from datetime import datetime
from typing import Optional
import openpyxl
from openpyxl.cell.cell import MergedCell


def _clean(value) -> str:
    """Convert a cell value to a clean single-line string.
    Returns '' if the value is blank/NaN."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    s = str(value).strip()
    if s.lower() == "nan" or s == "":
        return ""
    # Collapse internal newlines / tabs into a single space
    s = " ".join(s.splitlines())
    return " ".join(s.split())   # also collapse multiple spaces


class ExcelProcessor:
    def __init__(self, excel_path: str, sheet_name: str = "總表"):
        self.excel_path = excel_path
        self.sheet_name = sheet_name
        self.df: Optional[pd.DataFrame] = None
        self._original_columns = []

    def load(self) -> pd.DataFrame:
        self.df = pd.read_excel(self.excel_path, sheet_name=self.sheet_name)
        self._original_columns = list(self.df.columns)
        self.df["Unnamed: 0"] = self.df["Unnamed: 0"].ffill()
        return self.df

    def get_dataframe(self) -> pd.DataFrame:
        if self.df is None:
            self.load()
        return self.df

    def get_sheet_names(self) -> list:
        wb = openpyxl.load_workbook(self.excel_path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names

    def get_column_names(self) -> list:
        if self.df is None:
            self.load()
        return list(self.df.columns)

    def get_display_headers(self) -> list:
        if self.df is None:
            self.load()
        headers = []
        for col in self.df.columns:
            if isinstance(col, datetime):
                headers.append(col.strftime("%Y/%m/%d"))
            else:
                headers.append(str(col))
        return headers

    def get_display_data(self) -> list:
        if self.df is None:
            self.load()
        data = []
        for _, row in self.df.iterrows():
            row_data = []
            for val in row:
                if pd.isna(val):
                    row_data.append("")
                else:
                    row_data.append(str(val))
            data.append(row_data)
        return data

    def find_date_column(self, target_date: datetime) -> Optional[str]:
        if self.df is None:
            self.load()
        for col in self.df.columns:
            if isinstance(col, datetime) and col.date() == target_date.date():
                return col
            elif isinstance(col, str):
                if (target_date.strftime("%Y/%m/%d") in col or
                    target_date.strftime("%#m/%#d") in col or
                    target_date.strftime("%m/%d") in col):
                    return col
        return None

    def extract_tasks(self, target_date: datetime):
        if self.df is None:
            self.load()

        date_column = self.find_date_column(target_date)
        if date_column is None:
            raise ValueError(
                f"找不到對應的日期欄位：{target_date.strftime('%Y-%m-%d')}\n"
                f"Excel 中的欄位：{list(self.df.columns)}"
            )

        shipyard_tasks = []
        engine_room_tasks = []

        for _, row in self.df.iterrows():
            category = _clean(row.iloc[0])
            code     = _clean(row.iloc[1])
            name     = _clean(row.iloc[2])
            progress = _clean(row[date_column])

            # Skip if progress cell is blank
            if not progress:
                continue

            if category == "船廠工程":
                line = f"{code} - {progress}" if code else progress
                shipyard_tasks.append((name, line))
            elif "自修" in category:
                line = f"{name} - {progress}" if name else progress
                engine_room_tasks.append((name, line))

        return shipyard_tasks, engine_room_tasks

    def update_cell(self, row: int, col: int, value) -> None:
        if self.df is None:
            self.load()
        self.df.iat[row, col] = value

    def save_from_table_data(self, data: list, headers: list, output_path: str = None):
        """Save table data back to Excel, only updating cell values in-place.
        Preserves all formatting (colors, fonts, borders, merges, etc.)."""
        path = output_path or self.excel_path

        wb = openpyxl.load_workbook(path)
        ws = wb[self.sheet_name]

        # Row 1 is the header row, data starts at row 2
        # Only update data cells, leave headers untouched
        # Skip merged cells (read-only) — only the top-left cell of a merge is writable
        for r_idx, row_data in enumerate(data):
            for c_idx, value in enumerate(row_data):
                cell = ws.cell(row=r_idx + 2, column=c_idx + 1)
                if isinstance(cell, MergedCell):
                    continue
                # Convert empty strings back to None to preserve blank cells
                if value == "":
                    cell.value = None
                else:
                    # Try to preserve numeric types
                    try:
                        cell.value = float(value)
                    except (ValueError, TypeError):
                        cell.value = value

        wb.save(path)
        wb.close()
