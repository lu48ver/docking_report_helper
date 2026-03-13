from datetime import datetime
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _one_line(text: str) -> str:
    """Collapse any newlines / extra spaces in a string to a single line."""
    return " ".join(text.split())


class ReportGenerator:
    def __init__(self, template_path: str):
        self.template_path = template_path
        self.doc = None

    # Department mapping: key -> (report_title, english_marker)
    DEPT_MAP = {
        "engine": ("機艙塢修報告", "Eeg Dept."),
        "deck":   ("甲板塢修報告", "Deck Dept."),
        "none":   ("塢修報告", ""),
    }

    def generate(
        self,
        target_date: datetime,
        start_date: datetime,
        shipyard_tasks: list,
        engine_room_tasks: list,
        output_path: str,
        table_layout: str = "span",
        department: str = "engine",
        weather_text: str = "",
    ) -> str:
        self.doc = Document(self.template_path)

        day_number = (target_date - start_date).days
        day_str = f"D{day_number}"
        day_date_str = target_date.strftime("%Y%m%d")
        _MONTHS_EN = [
            "January", "February", "March", "April", "May", "June",
            "July", "August", "September", "October", "November", "December",
        ]
        formatted_date_str = (
            f"{target_date.day:02d}-{_MONTHS_EN[target_date.month - 1]}-{target_date.year}"
        )

        dept_title, dept_eng = self.DEPT_MAP.get(department, self.DEPT_MAP["engine"])

        # Replace date placeholders
        self._replace_placeholders(
            target_date, day_str, day_date_str, formatted_date_str,
            dept_title=dept_title, dept_eng=dept_eng, weather_text=weather_text,
        )

        # Filter new tasks (not already in template)
        new_shipyard = self._filter_new_tasks(shipyard_tasks)
        new_engine = self._filter_new_tasks(engine_room_tasks)

        # Insert task paragraphs — match both old and new anchor text
        self._insert_task_paragraphs("船廠本日工作項目:", shipyard_tasks)
        self._insert_task_paragraphs("自修工作", engine_room_tasks)

        # Insert task tables after last task
        self._insert_task_table(new_shipyard, table_layout=table_layout)
        self._insert_task_table(new_engine, table_layout=table_layout)

        self.doc.save(output_path)
        return output_path

    def _replace_placeholders(
        self, target_date, day_str, day_date_str, formatted_date_str,
        dept_title="機艙塢修報告", dept_eng="Eeg Dept.", weather_text="",
    ):
        weather_inserted = False
        for i, para in enumerate(self.doc.paragraphs):
            if "DATE：" in para.text:
                para.text = f"DATE：{target_date.strftime('%Y/%m/%d')}"
            elif "塢修報告" in para.text and ("Dept." in para.text or "Day-" in para.text):
                # Build title line: e.g. "機艙塢修報告Eeg Dept.-D4-20260308"
                if dept_eng:
                    para.text = f"{dept_title}{dept_eng}-{day_str}-{day_date_str} "
                else:
                    para.text = f"{dept_title}-{day_str}-{day_date_str} "

                # Insert weather line right after this paragraph
                if weather_text and not weather_inserted:
                    weather_para = Paragraph(OxmlElement("w:p"), para._parent)
                    weather_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    run = weather_para.add_run(weather_text)
                    run.font.size = Pt(10)
                    para._element.addnext(weather_para._element)
                    weather_inserted = True
            elif "Date:" in para.text:
                parts = para.text.split("Date:")
                if len(parts) == 2:
                    prefix = parts[0].strip()
                    para.text = f"{prefix} Date: {formatted_date_str}"

    def _insert_task_paragraphs(self, anchor_text: str, task_list: list):
        for para in self.doc.paragraphs:
            # Match anchor_text and also backward-compatible "機艙自修工作"
            if anchor_text in para.text or (anchor_text == "自修工作" and "機艙自修工作" in para.text):
                new_paragraphs = []
                for idx, (_, line) in enumerate(task_list, 1):
                    new_para = Paragraph(OxmlElement("w:p"), para._parent)
                    new_para.add_run(f"{idx}. {_one_line(line)}")
                    new_paragraphs.append(new_para)

                for new_para in reversed(new_paragraphs):
                    para._element.addnext(new_para._element)
                break

    def _insert_task_table(self, task_list: list, table_layout: str = "span"):
        if not task_list:
            return

        if table_layout == "span":
            self._insert_table_span(task_list)
        elif table_layout == "single":
            self._insert_table_single(task_list)
        else:
            self._insert_table_grid(task_list)

    def _find_last_task_paragraph(self, task_list):
        """Find the paragraph of the last numbered task item."""
        last_line = f"{len(task_list)}. {task_list[-1][1]}"
        for para in self.doc.paragraphs:
            if para.text.strip() == last_line.strip():
                return para
        return None

    @staticmethod
    def _set_row_height(row, height_cm):
        """Set a table row to an exact height."""
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        trHeight = OxmlElement("w:trHeight")
        trHeight.set(qn("w:val"), str(int(height_cm * 567)))  # 1 cm ≈ 567 twips
        trHeight.set(qn("w:hRule"), "exact")
        trPr.append(trHeight)

    @staticmethod
    def _center_cell_text(cell):
        """Center-align all paragraphs in a cell."""
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── Layout A: 跨欄置中 ──────────────────────────────────
    def _insert_table_span(self, task_list):
        anchor = self._find_last_task_paragraph(task_list)
        if not anchor:
            return

        total = len(task_list)
        # Each task = 1 merged header row + 1 photo row (2 cols)
        table = self.doc.add_table(rows=total * 2, cols=2)
        table.style = "Table Grid"

        for idx, (_, line) in enumerate(task_list):
            header_row_idx = idx * 2
            photo_row_idx = idx * 2 + 1

            # Merge header row across 2 columns
            cell_a = table.cell(header_row_idx, 0)
            cell_b = table.cell(header_row_idx, 1)
            merged = cell_a.merge(cell_b)
            merged.text = _one_line(line)
            self._center_cell_text(merged)

            # Set photo row height to ~3cm for pasting photos
            self._set_row_height(table.rows[photo_row_idx], 3.0)

        anchor._element.addnext(table._element)

    # ── Layout B: 單欄排列 ──────────────────────────────────
    def _insert_table_single(self, task_list):
        anchor = self._find_last_task_paragraph(task_list)
        if not anchor:
            return

        total = len(task_list)
        # Each task = 1 content row + 1 blank photo row
        table = self.doc.add_table(rows=total * 2, cols=1)
        table.style = "Table Grid"

        for idx, (_, line) in enumerate(task_list):
            content_row_idx = idx * 2
            photo_row_idx = idx * 2 + 1

            table.cell(content_row_idx, 0).text = _one_line(line)

            # Set photo row height to ~3cm
            self._set_row_height(table.rows[photo_row_idx], 3.0)

        anchor._element.addnext(table._element)

    # ── Layout C: 雙欄並排（原有邏輯）──────────────────────
    def _insert_table_grid(self, task_list, cols=2):
        anchor = self._find_last_task_paragraph(task_list)
        if not anchor:
            return

        total = len(task_list)
        rows = (total + cols - 1) // cols
        table = self.doc.add_table(rows=rows * 2, cols=cols)
        table.style = "Table Grid"

        for idx, (_, line) in enumerate(task_list):
            row = (idx // cols) * 2
            col = idx % cols
            table.cell(row, col).text = _one_line(line)

        anchor._element.addnext(table._element)

    def _filter_new_tasks(self, task_list: list) -> list:
        existing_lines = [para.text.strip() for para in self.doc.paragraphs]
        new_tasks = []
        for name, line in task_list:
            if f"- {line}" not in existing_lines and line not in existing_lines:
                new_tasks.append((name, line))
        return new_tasks
