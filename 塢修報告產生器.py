
import os
import pandas as pd
from datetime import datetime
from docx import Document
import tkinter as tk
from tkinter import filedialog, messagebox

# ==== GUI 設定 ====
root = tk.Tk()
root.withdraw()
import json
config_path = "last_paths.json"

# === 讀取記憶設定 ===
last_excel = None
last_template = None
last_output_dir = None

if os.path.exists(config_path):
    with open(config_path, "r", encoding="utf-8") as f:
        last_data = json.load(f)
        last_excel = last_data.get("excel_path")
        last_template = last_data.get("template_path")
        last_output_dir = last_data.get("output_dir")

# Excel 檔選擇 + 記憶功能
last_excel = None
if os.path.exists(config_path):
    with open(config_path, "r", encoding="utf-8") as f:
        last_data = json.load(f)
        last_excel = last_data.get("excel_path")

    if last_excel and os.path.exists(last_excel):
        use_last = messagebox.askyesno("使用上次的 Excel 檔案？", f"上次選擇的 Excel 檔案：\n{last_excel}\n\n是否使用這個檔案？")
        if not use_last:
            last_excel = None

if not last_excel:
    last_excel = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
    if not last_excel:
        raise ValueError("未選擇 Excel 檔案，程式結束")

excel_path = last_excel
# ==== 基本參數設定 ====
from datetime import datetime
from tkinter import simpledialog

# ==== 日期自動帶入 + 可選擇 ====
today = datetime.today()
today_str = today.strftime("%Y-%m-%d")

# 詢問是否使用今天日期
use_today = messagebox.askyesno("選擇日期", f"是否使用今天日期？\n（{today_str}）")

if use_today:
    target_date = today
else:
    # 使用者輸入日期（格式：YYYY-MM-DD）
    input_date_str = simpledialog.askstring("輸入日期", "請輸入日期（格式：YYYY-MM-DD）：", initialvalue=today_str)
    if not input_date_str:
        raise ValueError("未輸入日期，程式結束")
    try:
        target_date = datetime.strptime(input_date_str, "%Y-%m-%d")
    except ValueError:
        raise ValueError("日期格式錯誤，請輸入 YYYY-MM-DD")

# ==== 其餘變數照常計算 ====
target_date_str = target_date.strftime("%Y-%m-%d")
start_date = datetime.strptime("2025-04-24", "%Y-%m-%d")  # 根據你實際定義的起始日
day_number = (target_date - start_date).days
day_str = f"D{day_number}"
day_date_str = target_date.strftime("%Y%m%d")
formatted_date_str = target_date.strftime("%d-%B-%Y")
# 選擇 Word 樣板
last_template = None
if os.path.exists(config_path):
    with open(config_path, "r", encoding="utf-8") as f:
        last_data = json.load(f)
        last_template = last_data.get("template_path")

    if last_template and os.path.exists(last_template):
        use_last = messagebox.askyesno("使用上次的 Word 樣板？", f"上次選擇的 Word 樣板：\n{last_template}\n\n是否使用這個檔案？")
        if not use_last:
            last_template = None

if not last_template:
    last_template = filedialog.askopenfilename(title="請選擇 Word 樣板檔案", filetypes=[("Word files", "*.docx")])
    if not last_template:
        raise ValueError("未選擇 Word 樣板，程式結束")

template_path = last_template

# 輸出路徑
if last_output_dir and os.path.exists(last_output_dir):
    use_last = messagebox.askyesno("使用上次的輸出資料夾？", f"上次選擇的輸出資料夾：\n{last_output_dir}\n\n是否使用這個資料夾？")
    if not use_last:
        last_output_dir = None

if not last_output_dir:
    messagebox.showinfo("選擇輸出資料夾", "請選擇要儲存報告與照片的資料夾")
    last_output_dir = filedialog.askdirectory()
    if not last_output_dir:
        raise ValueError("未選擇輸出資料夾，程式結束")

output_dir = last_output_dir
output_path = os.path.join(output_dir, f"SY-機艙塢修報告-{day_date_str}-{day_str}.docx")
photo_root = output_dir

# ==== 載入資料 ====
df = pd.read_excel(excel_path, sheet_name="總表")
df['Unnamed: 0'].fillna(method='ffill', inplace=True)

print("🔍 Excel 欄位名稱如下：")
print(df.columns.tolist())

# ==== 嘗試抓取日期欄位 ====
date_column = None
for col in df.columns:
    if isinstance(col, datetime) and col.date() == target_date.date():
        date_column = col
        break
    elif isinstance(col, str) and (
        target_date.strftime("%Y/%m/%d") in col or
        target_date.strftime("%#m/%#d") in col or
        target_date.strftime("%m/%d") in col
    ):
        date_column = col
        break

if not date_column:
    print("❌ 找不到對應的日期欄位！你設定的是：", target_date_str)
    print("🧾 Excel 中的欄位如下：", df.columns.tolist())
    raise ValueError("無法找到對應日期欄")

# ==== 過濾工作項目 ====
shipyard_tasks = []
engine_room_tasks = []

for _, row in df.iterrows():
    category = str(row['Unnamed: 0']).strip()
    code = str(row['Unnamed: 1']).strip()
    name = str(row['Unnamed: 2']).strip()
    progress = row[date_column]

    if pd.notna(progress) and progress != "nan":
        if category == "船廠工程":
            line = f"{code} - {progress}"
            shipyard_tasks.append((name, line))
        elif category == "機艙自修":
            line = f"{name} - {progress}"
            engine_room_tasks.append((name, line))

print(f"抓到船廠項目 {len(shipyard_tasks)} 筆")
print(f"抓到機艙項目 {len(engine_room_tasks)} 筆")
messagebox.showinfo("資料載入結果", f"抓到船廠項目：{len(shipyard_tasks)} 筆\n抓到機艙項目：{len(engine_room_tasks)} 筆")

# ==== 編輯 Word 模板 ====
doc = Document(template_path)
for para in doc.paragraphs:
    if "DATE：" in para.text:
        para.text = f"DATE：{target_date.strftime('%Y/%m/%d')}"
    elif "塢修報告Eeg Dept.-Day-" in para.text:
        para.text = f"塢修報告Eeg Dept.-{day_str}-{day_date_str} "
    elif "Date:" in para.text:
        parts = para.text.split("Date:")
        if len(parts) == 2:
            prefix = parts[0].strip()
            para.text = f"{prefix} Date: {formatted_date_str}"


from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

def restore_task_paragraphs(doc, anchor_text, task_list):
    """
    在指定 anchor_text 段落下，依順序插入「1. 任務內容」，保持正確順序。
    """
    for para in doc.paragraphs:
        if anchor_text in para.text:
            new_paragraphs = []
            for idx, (_, line) in enumerate(task_list, 1):
                new_para = Paragraph(OxmlElement("w:p"), para._parent)
                new_para.add_run(f"{idx}. {line}")
                new_paragraphs.append(new_para)

            # 反向插入，讓順序正確
            for new_para in reversed(new_paragraphs):
                para._element.addnext(new_para._element)
            break
def insert_task_table_after_last_task(doc, task_list, cols=2):
    """
    在清單段落（最後一筆）之後插入表格，列出進度內容，第二列保留空白。
    """
    if not task_list:
        return

    # 找最後一筆清單文字內容
    last_line = f"{len(task_list)}. {task_list[-1][1]}"
    
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() == last_line.strip():
            total = len(task_list)
            rows = (total + cols - 1) // cols
            table = doc.add_table(rows=rows * 2, cols=cols)
            table.style = 'Table Grid'

            for idx, (_, line) in enumerate(task_list):
                row = (idx // cols) * 2
                col = idx % cols
                table.cell(row, col).text = line.strip()

            # 插入在最後清單段落之後
            para._element.addnext(table._element)
            break

def filter_new_tasks(doc, task_list):
    """
    檢查 Word 裡已經有的清單項目，過濾出新的項目（根據任務內容）。
    """
    existing_lines = [para.text.strip() for para in doc.paragraphs]
    new_tasks = []
    for name, line in task_list:
        if f"- {line}" not in existing_lines and f"{line}" not in existing_lines:
            new_tasks.append((name, line))
    return new_tasks
# 過濾新增項目
new_shipyard_tasks = filter_new_tasks(doc, shipyard_tasks)
new_engine_room_tasks = filter_new_tasks(doc, engine_room_tasks)
restore_task_paragraphs(doc, "船廠本日工作項目:", shipyard_tasks)
restore_task_paragraphs(doc, "機艙自修工作", engine_room_tasks)

insert_task_table_after_last_task(doc, new_shipyard_tasks)
insert_task_table_after_last_task(doc, new_engine_room_tasks)
# 儲存報告
doc.save(output_path)

# 建立照片資料夾
photo_day_folder = os.path.join(photo_root, day_date_str)
os.makedirs(photo_day_folder, exist_ok=True)

# 只為新任務建立資料夾（不重複建立已存在的）
# 為船廠工作建立子資料夾（不加日期）
for _, line in new_shipyard_tasks:
    if line and line != "nan":
        folder_path = os.path.join(photo_day_folder, line.strip())
        if not os.path.exists(folder_path):
            os.makedirs(folder_path)

# 為機艙自修工作建立子資料夾（前綴加上日期）
for _, line in new_engine_room_tasks:
    if line and line != "nan":
        folder_name = f"{day_date_str}_{line.strip()}"
        folder_path = os.path.join(photo_day_folder, folder_name)
        if not os.path.exists(folder_path):
            os.makedirs(folder_path)

messagebox.showinfo("完成", f"報告已儲存至：\n{output_path}\n\n照片資料夾也已建立。")

# ✅ 儲存 Excel/Word 路徑記憶檔
with open(config_path, "w", encoding="utf-8") as f:
    json.dump({
        "excel_path": excel_path,
        "template_path": template_path,
        "output_dir": output_dir
    }, f, ensure_ascii=False, indent=2)